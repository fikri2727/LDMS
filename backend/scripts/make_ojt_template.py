"""Builds the user-friendly OJT upload template.

The answer cells MUST stay where backend/app/services/excel.py reads them:
  C3 Title, C4 Venue, C5 Trainer Type, C6 Trainer Name, C7 Start Date, C8 End Date,
  C9 Start Time, C10 End Time; participants' Staff No in column B from row 13.
"""

import sys
from datetime import datetime, time

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
ROWS = 100  # participant rows (13 .. 112)

GREEN = "1D8E72"
GREEN_DARK = "15664F"
GREEN_LIGHT = "E6F4F0"
INPUT = "FFF8DC"  # pale yellow = "type here"
GREY_TEXT = "6B7280"
BORDER = "C9CED6"

thin = Side(style="thin", color=BORDER)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
FONT = "Calibri"

wb = Workbook()
ws = wb.active
ws.title = "OJT Upload"
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 100

for col, width in {"A": 6, "B": 22, "C": 30, "D": 12, "E": 12, "F": 44}.items():
    ws.column_dimensions[col].width = width

# ---- Row 1-2: header + one-line instructions ----
ws.merge_cells("A1:F1")
ws["A1"] = "TAMCO LDMS  ·  OJT Upload Template"
ws["A1"].font = Font(name=FONT, size=16, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor=GREEN)
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 32

ws.merge_cells("A2:F2")
ws["A2"] = (
    "Fill in the yellow boxes, list the participants' Staff No below, save, then upload in "
    "LDMS → Training Records → OJT → Upload Excel.   ✱ = required"
)
ws["A2"].font = Font(name=FONT, size=10, italic=True, color=GREEN_DARK)
ws["A2"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
ws["A2"].alignment = Alignment(vertical="center", wrap_text=True, indent=1)
ws.row_dimensions[2].height = 30

# ---- Rows 3-10: OJT details (label A:B, answer C:E merged, tip F) ----
details = [
    (3, "Title ✱", "OJT / training title", None),
    (4, "Venue ✱", "Where the OJT took place, e.g. PLANT 2 - LINE 3", None),
    (5, "Trainer Type ✱", "Choose from the list: Internal or External", None),
    (6, "Trainer Name ✱", "Internal: the staff member's name or Staff No\nExternal: the trainer's / company's name", None),
    (7, "Start Date ✱", "dd/mm/yyyy, e.g. 21/10/2026", "date"),
    (8, "End Date ✱", "dd/mm/yyyy - same as Start Date for a 1-day OJT", "date"),
    (9, "Start Time ✱", "e.g. 8:30 AM", "time"),
    (10, "End Time ✱", "e.g. 5:00 PM", "time"),
]
for row, label, tip, kind in details:
    ws.merge_cells(f"A{row}:B{row}")
    ws.merge_cells(f"C{row}:E{row}")
    a = ws[f"A{row}"]
    a.value = label
    a.font = Font(name=FONT, size=11, bold=True, color="1F2937")
    a.fill = PatternFill("solid", fgColor="F3F4F6")
    a.alignment = Alignment(vertical="center", indent=1)
    c = ws[f"C{row}"]
    c.fill = PatternFill("solid", fgColor=INPUT)
    c.font = Font(name=FONT, size=11, bold=True)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    if kind == "date":
        c.number_format = "dd/mm/yyyy"
    elif kind == "time":
        c.number_format = "h:mm AM/PM"
    for col in "ABCDE":
        ws[f"{col}{row}"].border = box
    f = ws[f"F{row}"]
    f.value = tip
    f.font = Font(name=FONT, size=9, italic=True, color=GREY_TEXT)
    f.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[row].height = 30 if "\n" in tip else 22

# Trainer Type drop-down
dv_type = DataValidation(type="list", formula1='"Internal,External"', allow_blank=True)
dv_type.error, dv_type.errorTitle = "Please choose Internal or External.", "Trainer Type"
dv_type.prompt, dv_type.promptTitle = "Choose Internal or External", "Trainer Type"
dv_type.add("C5")
# Dates / times: Excel warns on obviously wrong entries (text like 21/10/2026 still works when typed as a date)
dv_date = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2035,12,31)", allow_blank=True)
dv_date.error, dv_date.errorTitle = "Enter a date as dd/mm/yyyy, e.g. 21/10/2026.", "Date"
dv_date.prompt, dv_date.promptTitle = "dd/mm/yyyy", "Date"
dv_date.add("C7")
dv_date.add("C8")
dv_time = DataValidation(type="time", operator="between", formula1="TIME(0,0,0)", formula2="TIME(23,59,59)", allow_blank=True)
dv_time.error, dv_time.errorTitle = "Enter a time, e.g. 8:30 AM or 17:00.", "Time"
dv_time.prompt, dv_time.promptTitle = "e.g. 8:30 AM", "Time"
dv_time.add("C9")
dv_time.add("C10")
for dv in (dv_type, dv_date, dv_time):
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)

# ---- Row 11: section title; Row 12: participant header ----
ws.merge_cells("A11:F11")
ws["A11"] = "PARTICIPANTS  -  one Staff No per row (leading zeros are kept, e.g. 001622)"
ws["A11"].font = Font(name=FONT, size=11, bold=True, color=GREEN_DARK)
ws["A11"].alignment = Alignment(vertical="bottom", indent=0)
ws.row_dimensions[11].height = 26

headers = {"A": "No", "B": "Staff No ✱", "C": "Name (optional - for your reference)"}
ws.merge_cells("C12:F12")
for col, text in headers.items():
    cell = ws[f"{col}12"]
    cell.value = text
    cell.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=GREEN)
    cell.alignment = Alignment(horizontal="center" if col == "A" else "left", vertical="center", indent=0 if col == "A" else 1)
for col in "ABCDEF":
    ws[f"{col}12"].border = box
    ws[f"{col}12"].fill = PatternFill("solid", fgColor=GREEN)
ws.row_dimensions[12].height = 22
ws["B12"].comment = Comment("Type the Staff No exactly as in LDMS, e.g. 001622.\nThis column is formatted as text so leading zeros stay.", "LDMS")

stripe = PatternFill("solid", fgColor="F9FAFB")
for i in range(ROWS):
    r = 13 + i
    ws.merge_cells(f"C{r}:F{r}")
    n = ws[f"A{r}"]
    n.value = i + 1
    n.font = Font(name=FONT, size=10, color=GREY_TEXT)
    n.alignment = Alignment(horizontal="center", vertical="center")
    s = ws[f"B{r}"]
    s.number_format = "@"  # text: keeps 001622
    s.fill = PatternFill("solid", fgColor=INPUT)
    s.font = Font(name=FONT, size=11, bold=True)
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws[f"C{r}"].font = Font(name=FONT, size=10, color=GREY_TEXT)
    ws[f"C{r}"].alignment = Alignment(vertical="center", indent=1)
    for col in "ABCDEF":
        ws[f"{col}{r}"].border = box
        if col != "B" and i % 2 == 1:
            ws[f"{col}{r}"].fill = stripe

ws.freeze_panes = "A13"  # header + details stay visible while scrolling the participant list
ws.print_title_rows = "12:12"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_area = f"A1:F{12 + ROWS}"

# ---- Sheet 2: how to use + example ----
ex = wb.create_sheet("How to use")
ex.sheet_view.showGridLines = False
ex.column_dimensions["A"].width = 4
ex.column_dimensions["B"].width = 26
ex.column_dimensions["C"].width = 70
ex.merge_cells("B1:C1")
ex["B1"] = "How to use the OJT upload template"
ex["B1"].font = Font(name=FONT, size=16, bold=True, color=GREEN_DARK)
ex.row_dimensions[1].height = 30
steps = [
    "Go to the 'OJT Upload' sheet (the first sheet) - only that sheet is read.",
    "Fill in every yellow box marked ✱ (Title, Venue, Trainer Type, Trainer Name, dates and times).",
    "Trainer Type: pick Internal (a TAMCO staff member) or External (an outside trainer / company).",
    "List each participant's Staff No in the yellow column, one per row, e.g. 001622. Names are optional.",
    "Don't move, insert or delete rows/columns - the system reads the boxes by position.",
    "Save the file (.xlsx) and upload it: LDMS > Training Records > OJT > Upload Excel.",
    "If something is wrong, LDMS tells you which box to fix (e.g. a Staff No that doesn't exist) and nothing is saved.",
]
for i, text in enumerate(steps, start=3):
    ex[f"B{i}"] = f"Step {i - 2}"
    ex[f"B{i}"].font = Font(name=FONT, size=11, bold=True, color=GREEN)
    ex[f"C{i}"] = text
    ex[f"C{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ex.row_dimensions[i].height = 30

row = len(steps) + 4
ex[f"B{row}"] = "Example"
ex[f"B{row}"].font = Font(name=FONT, size=13, bold=True, color=GREEN_DARK)
example = [
    ("Title", "FORKLIFT SAFE OPERATION"),
    ("Venue", "WAREHOUSE A"),
    ("Trainer Type", "Internal"),
    ("Trainer Name", "AHMAD BIN ALI"),
    ("Start Date", datetime(2026, 10, 21)),
    ("End Date", datetime(2026, 10, 21)),
    ("Start Time", time(8, 30)),
    ("End Time", time(12, 30)),
    ("Staff No (participants)", "001622, 001553, 000124 ..."),
]
for j, (k, v) in enumerate(example, start=row + 1):
    ex[f"B{j}"] = k
    ex[f"B{j}"].font = Font(name=FONT, size=10, bold=True)
    ex[f"B{j}"].fill = PatternFill("solid", fgColor="F3F4F6")
    ex[f"C{j}"] = v
    ex[f"C{j}"].fill = PatternFill("solid", fgColor=INPUT)
    if isinstance(v, datetime):
        ex[f"C{j}"].number_format = "dd/mm/yyyy"
    elif isinstance(v, time):
        ex[f"C{j}"].number_format = "h:mm AM/PM"
    ex[f"C{j}"].alignment = Alignment(horizontal="left")
    for col in "BC":
        ex[f"{col}{j}"].border = box

wb.active = 0
wb.save(OUT)
print("saved", OUT)
