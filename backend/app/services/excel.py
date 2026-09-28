"""Excel helpers (openpyxl) — ported from ldms-web/src/lib/staff-excel.ts and ojt-excel.ts."""

import io
import re
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.forms import bad


def open_first_sheet(data: bytes) -> Worksheet:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)  # data_only: formula cells give their cached result
    except Exception:
        raise bad("Could not read the uploaded file. Please upload an .xlsx file based on the provided template.")
    if not wb.worksheets:
        raise bad("The uploaded file has no sheets. Please use the provided template.")
    return wb.worksheets[0]


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(timespec="milliseconds") + "Z"
    if isinstance(v, date):
        return v.isoformat() + "T00:00:00.000Z"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def normalize_staff_no(raw: str) -> str:
    """Excel drops leading zeros ("001622" -> 1622); pad all-digit values back to 6 digits."""
    v = raw.strip().upper()
    return v.zfill(6) if re.fullmatch(r"\d{1,5}", v) else v


STAFF_HEADER_ALIASES: dict[str, list[str]] = {
    "staffNo": ["staff no", "staff no.", "staffno"],
    "staffName": ["staff name", "full name", "name"],
    "email": ["email"],
    "gender": ["gender"],
    "designation": ["designation"],
    "nationality": ["nationality"],
    "division": ["division"],
    "department": ["department"],
    "section": ["section"],
    "supervisorStaffNo": ["supervisor staff no", "supervisor staff no.", "supervisor"],
    "roleType": ["system role", "role"],
    "status": ["status"],
}


def parse_staff_excel(data: bytes) -> list[dict]:
    """Plain table, header row 1, one staff member per row. Headers matched by name (case-insensitive)."""
    ws = open_first_sheet(data)
    col: dict[str, int] = {}
    for c in ws[1]:
        text = cell_text(c.value).lower()
        for fld, aliases in STAFF_HEADER_ALIASES.items():
            if text in aliases:
                col[fld] = c.column
    if "staffNo" not in col:
        raise bad('Could not find a "Staff No" column in row 1. Please use the provided template.')

    rows = []
    for r in range(2, ws.max_row + 1):
        def f(key: str) -> str:
            return cell_text(ws.cell(r, col[key]).value) if key in col else ""

        staff_no = normalize_staff_no(f("staffNo"))
        if not staff_no:
            continue
        row = {k: f(k) for k in STAFF_HEADER_ALIASES}
        row["staffNo"] = staff_no
        row["supervisorStaffNo"] = normalize_staff_no(row["supervisorStaffNo"])
        row["rowNumber"] = r
        rows.append(row)

    if not rows:
        raise bad("No staff rows found. Add at least one row below the header, with a Staff No in each.")
    return rows


# ---------- OJT upload ----------
# Matches the company's "template upload OJT.xlsx": detail values in column C (rows 3-10),
# participant Staff Nos in column B from row 13.

OJT_CELL = {
    "title": "C3",
    "venue": "C4",
    "trainerType": "C5",
    "trainerName": "C6",
    "startDate": "C7",
    "endDate": "C8",
    "startTime": "C9",
    "endTime": "C10",
}
FIRST_PARTICIPANT_ROW = 13
PARTICIPANT_COLUMN = "B"


def _excel_serial(v) -> datetime | None:
    """A cell is either a real date/time (formatted) or a bare Excel serial number."""
    from datetime import time, timedelta

    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, time):
        return datetime(1899, 12, 30, v.hour, v.minute, v.second)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return datetime(1899, 12, 30) + timedelta(milliseconds=round(v * 86400000))
    return None


def _parse_date_cell(v, label: str) -> datetime:
    d = _excel_serial(v)
    if d:
        return datetime(d.year, d.month, d.day)
    text = cell_text(v)
    m = re.fullmatch(r"(\d{1,2})[/\-](\d{1,2})[/\-](\d{4})", text)
    if not m:
        raise bad(f'{label} must be in dd/mm/yyyy format, got "{text}".')
    try:
        return datetime(int(m[3]), int(m[2]), int(m[1]))
    except ValueError:
        raise bad(f'{label}: invalid date "{text}".')


def _parse_time_cell(v, label: str) -> str:
    d = _excel_serial(v)
    if d:
        return f"{d.hour:02d}:{d.minute:02d}"
    text = cell_text(v).upper()
    m = re.fullmatch(r"(\d{1,2}):(\d{2})\s*(AM|PM)?", text)
    if not m:
        raise bad(f'{label} must be in hh:mm AM/PM format, got "{text}".')
    hours = int(m[1])
    if m[3] == "PM" and hours < 12:
        hours += 12
    if m[3] == "AM" and hours == 12:
        hours = 0
    return f"{hours:02d}:{m[2]}"


def parse_ojt_excel(data: bytes) -> dict:
    from app.labels import TRAINER_TYPE_LABELS, enum_from_label

    ws = open_first_sheet(data)
    c = {k: ws[addr].value for k, addr in OJT_CELL.items()}
    title = cell_text(c["title"]).upper()
    venue = cell_text(c["venue"]).upper()
    trainer_type_raw = cell_text(c["trainerType"]).upper()
    trainer_name = cell_text(c["trainerName"]).upper()
    if not title or not venue or not trainer_name:
        raise bad(
            f"Title ({OJT_CELL['title']}), Venue ({OJT_CELL['venue']}), and Trainer Name ({OJT_CELL['trainerName']}) are required."
        )
    trainer_type = enum_from_label(TRAINER_TYPE_LABELS, trainer_type_raw)
    if not trainer_type:
        raise bad(f'Trainer Type ({OJT_CELL["trainerType"]}) must be "Internal" or "External", got "{trainer_type_raw}".')

    fields = {
        "title": title,
        "venue": venue,
        "trainer_type": trainer_type,
        "trainer_name": trainer_name,
        "start_date": _parse_date_cell(c["startDate"], f"Start Date ({OJT_CELL['startDate']})"),
        "end_date": _parse_date_cell(c["endDate"], f"End Date ({OJT_CELL['endDate']})"),
        "start_time": _parse_time_cell(c["startTime"], f"Start Time ({OJT_CELL['startTime']})"),
        "end_time": _parse_time_cell(c["endTime"], f"End Time ({OJT_CELL['endTime']})"),
    }

    staff_nos: list[str] = []
    for r in range(FIRST_PARTICIPANT_ROW, max(ws.max_row, FIRST_PARTICIPANT_ROW) + 1):
        no = cell_text(ws[f"{PARTICIPANT_COLUMN}{r}"].value).upper()
        if no and no not in staff_nos:
            staff_nos.append(no)
    if not staff_nos:
        raise bad(
            f"No participants found — list each Staff No in column {PARTICIPANT_COLUMN} starting row {FIRST_PARTICIPANT_ROW}."
        )
    return {"fields": fields, "staffNos": staff_nos}
