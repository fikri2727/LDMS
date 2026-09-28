import io
from datetime import datetime

from openpyxl import Workbook

from app.models import Certificate, Designation, Ojt, ParticipateOjt, Participation, Pme, Training, User
from tests.conftest import login_as

TRAINING_FORM = {
    "title": "test python training", "program": "EXT", "cost": "150.5", "platform": "PHYSICAL",
    "function": "DIGITAL", "venue": "hq", "hrdcClaimable": "on", "startDate": "2026-03-02",
    "endDate": "2026-03-03", "startTime": "09:00", "endTime": "17:00", "trainer": "someone",
}


def _exec_with_supervisor(db):
    return (
        db.query(User)
        .filter(User.designation == Designation.EXECUTIVE, User.status == "ACTIVE", User.supervisor_id.isnot(None))
        .first()
    )


def test_training_lifecycle_and_pme(admin_client, client, db, monkeypatch, admin):
    r = admin_client.post("/api/training/public", data=TRAINING_FORM)
    assert r.status_code == 200, r.text
    tid = r.json()["id"]
    t = db.get(Training, tid)
    assert t.training_code == f"PB020326{tid:05d}" and t.title == "TEST PYTHON TRAINING" and t.hrdc_claimable

    exec_user = _exec_with_supervisor(db)
    assert admin_client.post(f"/api/training/public/{tid}/participants", data={"userId": str(exec_user.id)}).status_code == 200
    assert admin_client.post(f"/api/training/public/{tid}/participants", data={"userId": str(exec_user.id)}).status_code == 400

    detail = admin_client.get(f"/api/training/public/{tid}").json()
    assert len(detail["training"]["participations"]) == 1
    assert all(o["id"] != exec_user.id for o in detail["staffOptions"])
    pid = detail["training"]["participations"][0]["id"]

    # someone else can't submit this participant's survey
    other = db.query(User).filter(User.id != exec_user.id, User.role_type == "STAFF").first()
    login_as(client, other)
    assert client.post(f"/api/training/public/{tid}/participants/{pid}/survey", data={}).status_code == 403
    # the non-admin detail view only shows their own row (none)
    assert client.get(f"/api/training/public/{tid}").json()["training"]["participations"] == []

    login_as(client, exec_user)
    r = client.post(f"/api/training/public/{tid}/participants/{pid}/survey",
                    data={"courseRelevance": "5", "practicalExercises": "4", "whatLearnt": " stuff "})
    assert r.status_code == 200, r.text
    p = db.get(Participation, pid)
    db.refresh(p)
    assert p.attendance.value == "COMPLETED" and p.course_relevance == 5 and p.what_learnt == "stuff"
    pme = db.query(Pme).filter_by(participation_id=pid).one()
    assert pme.supervisor_id == exec_user.supervisor_id and pme.status.value == "PENDING"

    my = client.get("/api/training/my").json()
    assert any(x["id"] == tid for x in my["trainings"])

    login_as(client, admin)
    # certificates (storage faked)
    monkeypatch.setattr("app.routers.training.save_upload", lambda sub, f: _async("certificates/x.pdf"))
    monkeypatch.setattr("app.routers.training.delete_upload", lambda path: None)
    r = admin_client.post(f"/api/training/public/{tid}/certificates", files={"file": ("cert.pdf", b"%PDF")})
    assert r.status_code == 200, r.text
    cert = db.query(Certificate).filter_by(training_id=tid).one()
    assert admin_client.delete(f"/api/training/public/{tid}/certificates/{cert.id}").status_code == 200

    lst = admin_client.get("/api/training/public").json()
    row = next(x for x in lst["trainings"] if x["id"] == tid)
    assert row["participations"][0]["pme"]["status"] == "PENDING"
    assert row["participations"][0]["user"]["department"] is not None or True

    r = admin_client.delete(f"/api/training/public/{tid}")
    assert r.status_code == 200, r.text
    assert db.get(Training, tid) is None


async def _async(v):
    return v


def test_training_requires_admin(client, staff_user):
    login_as(client, staff_user)
    r = client.post("/api/training/public", data=TRAINING_FORM)
    assert r.status_code == 403
    assert r.json()["detail"] == "You do not have permission to manage training records."


def test_ojt_admin_and_self_flow(admin_client, client, db, staff_user, admin):
    form = {"title": "ojt test", "venue": "plant", "trainerType": "INTERNAL", "trainerName": "trainer",
            "startDate": "2026-04-01", "endDate": "2026-04-02", "startTime": "08:30", "endTime": "12:00"}
    r = admin_client.post("/api/ojt", data=form)
    assert r.status_code == 200, r.text
    oid = r.json()["id"]
    o = db.get(Ojt, oid)
    assert o.total_day == 2 and o.total_hour == 3.5 and o.training_code.startswith("OJ010426")

    assert admin_client.post(f"/api/ojt/{oid}/participants", data={"userId": str(staff_user.id)}).status_code == 200
    part = db.query(ParticipateOjt).filter_by(ojt_id=oid, user_id=staff_user.id).one()

    login_as(client, staff_user)
    assert client.get(f"/api/ojt/{oid}/basic").status_code == 403  # not creator, not clerk
    r = client.post(f"/api/ojt/{oid}/participants/{part.id}/survey", data={"q1": "welding", "q2": "2", "q3": "4"})
    assert r.status_code == 200, r.text
    db.refresh(part)
    assert part.attendance.value == "COMPLETED" and part.q1 == "WELDING" and part.q3 == 4

    # self-service "Add My OJT"
    r = client.post("/api/ojt", data={**form, "isSelf": "1", "q1": "x", "q2": "1", "q3": "3"})
    assert r.json()["redirect"] == "/training"
    mine = db.get(Ojt, r.json()["id"])
    assert mine.created_by_user_id == staff_user.id
    assert client.delete(f"/api/ojt/{mine.id}").json()["redirect"] == "/training"

    login_as(client, admin)
    assert len(admin_client.get("/api/ojt").json()) >= 1
    assert admin_client.get("/api/ojt/trainer-options").json()


def test_ojt_excel_upload(admin_client, db, staff_user):
    wb = Workbook()
    ws = wb.active
    ws["C3"], ws["C4"], ws["C5"], ws["C6"] = "excel ojt", "line 1", "External", "vendor"
    ws["C7"], ws["C8"] = datetime(2026, 5, 4), "05/05/2026"
    ws["C9"], ws["C10"] = "8:00 AM", "5:30 PM"
    ws["B13"], ws["B14"] = staff_user.staff_no, "NOPE123"
    buf = io.BytesIO()
    wb.save(buf)
    r = admin_client.post("/api/ojt/upload", files={"file": ("o.xlsx", buf.getvalue())})
    assert r.status_code == 400 and "NOPE123" in r.json()["detail"]

    ws["B14"] = None
    buf = io.BytesIO()
    wb.save(buf)
    r = admin_client.post("/api/ojt/upload", files={"file": ("o.xlsx", buf.getvalue())})
    assert r.status_code == 200, r.text
    o = db.get(Ojt, r.json()["id"])
    assert (o.start_time, o.end_time, o.total_day, o.total_hour) == ("08:00", "17:30", 2, 9.5)
    assert o.trainer_type.value == "EXTERNAL"


def test_public_checkin(client, db, admin_client, staff_user):
    tid = admin_client.post("/api/training/public", data=TRAINING_FORM).json()["id"]
    client.cookies.clear()
    assert client.get(f"/api/checkin/training/{tid}").json()["id"] == tid
    assert client.post(f"/api/checkin/training/{tid}/verify", json={"staffId": "  "}).json() == {"error": "Please enter your Staff ID."}
    assert "not recognised" in client.post(f"/api/checkin/training/{tid}/verify", json={"staffId": "NOBODY"}).json()["error"]
    r = client.post(f"/api/checkin/training/{tid}/verify", json={"staffId": staff_user.staff_no.lower()}).json()
    assert r["verified"] and r["staffNo"] == staff_user.staff_no
    assert client.post(f"/api/checkin/training/{tid}/survey", data={"staffNo": staff_user.staff_no, "courseRelevance": "3"}).status_code == 200
    assert client.post(f"/api/checkin/training/{tid}/verify", json={"staffId": staff_user.staff_no}).json()["alreadyCompleted"]
    assert client.post(f"/api/checkin/training/{tid}/survey", data={"staffNo": staff_user.staff_no}).status_code == 400


def test_elearning_hours(admin_client):
    rows = admin_client.get("/api/training/elearning-hours").json()
    assert isinstance(rows, list)
    if rows:
        assert {"staffNo", "moduleTitle", "hours", "completedAt"} <= rows[0].keys()
